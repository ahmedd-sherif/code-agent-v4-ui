from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def create_word_doc():
    doc = Document()
    
    # Set margins to be moderate
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style configuration
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Segoe UI'
    font.size = Pt(10)

    # Title
    heading = doc.add_paragraph('CSAI810 Project Definition')
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(16)

    # Header Info
    p = doc.add_paragraph()
    p.add_run('Group Category: CSAI810 Groups').bold = True
    p.add_run('\n')
    p.add_run('Group Name: CSAI 810 Section 44').bold = True
    
    # Read MD content manually to allow flexibility
    # Or just hardcode the structure since we have the content in memory
    
    # 1. Title
    h = doc.add_heading('1. Title of your project', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None # default color
    doc.add_paragraph('Autonomous Code Agent').runs[0].bold = True

    # 2. Goals
    h = doc.add_heading('2. Short description of goals and limitations', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None

    p = doc.add_paragraph()
    p.add_run('Goals: ').bold = True
    p.add_run('The main goal is to build a comprehensive, local-first AI coding assistant that can autonomously help developers build software. Unlike standard chatbots, the Autonomous Code Agent enables the LLM to actively interact with the development environment—writing files, executing terminal commands, managing git repositories, and performing web searches to solve complex coding tasks. It aims to bridge the gap between "chatting about code" and "doing the coding."')

    p = doc.add_paragraph()
    p.add_run('Limitations:').bold = True
    
    doc.add_paragraph('Context Window: The agent\'s understanding of large codebases is limited by the context window of the underlying models.', style='List Bullet')
    doc.add_paragraph('Visual Validation: The agent cannot "see" the UI it builds (no computer vision integration yet).', style='List Bullet')
    doc.add_paragraph('Model Dependency: Performance heavily depends on the reasoning capability of the selected model.', style='List Bullet')

    # 3. Achieve
    h = doc.add_heading('3. What will you achieve (learn + produce) and what is your prior knowledge?', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None

    p = doc.add_paragraph()
    p.add_run('Produce: ').bold = True
    p.add_run('We will produce a fully functional Streamlit-based application (deployable as a Local Web App or Desktop Interface) that serves as the interface for the agent. The system will include a persistent memory layer (SQLite), a robust toolset, and a flexible model switcher.')

    p = doc.add_paragraph()
    p.add_run('Learn: ').bold = True
    p.add_run('We will learn how to:')
    
    doc.add_paragraph('Design tool-use loops (ReAct pattern) for LLMs.', style='List Bullet')
    doc.add_paragraph('Manage structured output and function calling across different model providers.', style='List Bullet')
    doc.add_paragraph('Implement persistent conversation history and project context management.', style='List Bullet')
    doc.add_paragraph('Integrate local development tools (Git, Terminal) safely with an AI agent.', style='List Bullet')

    p = doc.add_paragraph()
    p.add_run('Prior Knowledge: ').bold = True
    p.add_run('We have prior experience with Python programming, API integrations (REST), and basic prompting.')

    # 4. Why
    h = doc.add_heading('4. Why do you want to do this project?', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None
    doc.add_paragraph('We want to move beyond simple "Chat Q&A" interfaces. Real developer productivity comes from an agent that can act—scaffold projects, fix bugs directly in files, and handle version control. This project allows us to explore the cutting edge of Agentic AI while building a tool that we can actually use in our daily workflow.')

    # 5. Deliverable
    h = doc.add_heading('5. Final deliverable and user', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None
    
    p = doc.add_paragraph()
    p.add_run('Final Deliverable: ').bold = True
    p.add_run('A software package (Python/Streamlit) containing:')
    doc.add_paragraph('The "Code Agent" application.', style='List Bullet')
    doc.add_paragraph('A suite of integrated tools (File System, Git, Terminal).', style='List Bullet')
    doc.add_paragraph('A unified UI for chat, workspace management, and model selection.', style='List Bullet')

    p = doc.add_paragraph()
    p.add_run('Target User: ').bold = True
    p.add_run('Software developers, students, and hobbyists who want an AI assistant that runs locally.')

    # 6. Boundaries
    h = doc.add_heading('6. Project boundaries (30 Hours Scope)', level=2)
    h.runs[0].font.size = Pt(12)
    h.runs[0].font.color.rgb = None

    p = doc.add_paragraph()
    p.add_run('Included (In-Scope):').bold = True
    doc.add_paragraph('Core Agent Loop (LiteLLM/Google GenAI)', style='List Bullet')
    doc.add_paragraph('Model Infrastructure: Integration with Ollama for local model inference and Ngrok.', style='List Bullet')
    doc.add_paragraph('Tool Implementation:', style='List Bullet')
    
    sub = doc.add_paragraph('File System (Read/Write/List)', style='List Bullet 2')
    sub.paragraph_format.left_indent = Inches(0.5)
    sub = doc.add_paragraph('Terminal (Run safe commands)', style='List Bullet 2')
    sub.paragraph_format.left_indent = Inches(0.5)
    sub = doc.add_paragraph('Git (Commit, Push, Pull)', style='List Bullet 2')
    sub.paragraph_format.left_indent = Inches(0.5)
    sub = doc.add_paragraph('Web Search (Tavily/DuckDuckGo)', style='List Bullet 2')
    sub.paragraph_format.left_indent = Inches(0.5)
    
    doc.add_paragraph('UI: Streamlit interface with sticky header.', style='List Bullet')
    doc.add_paragraph('Memory: Local SQLite database.', style='List Bullet')

    p = doc.add_paragraph()
    p.add_run('Excluded (Out-of-Scope):').bold = True
    doc.add_paragraph('Multi-User Auth', style='List Bullet')
    doc.add_paragraph('Visual UI Testing', style='List Bullet')
    doc.add_paragraph('IDE Extensions', style='List Bullet')
    doc.add_paragraph('Voice Mode', style='List Bullet')

    doc.save('definition/Project_Definition.docx')

if __name__ == "__main__":
    create_word_doc()
    print("Created definition/Project_Definition.docx")
